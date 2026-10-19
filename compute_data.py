import openpyxl, json, datetime
from collections import defaultdict

F1 = r"C:\Users\kirby\Desktop\简历\赛柯实习\ERP_hub_采购订单 ！ (1).xlsx"
F2 = r"C:\Users\kirby\Desktop\简历\赛柯实习\ERP_hub_应付.xlsx"
FX = {"USD": 7.2, "EUR": 7.8, "CNY": 1.0}

def cny(amount, ccy):
    try:
        return float(amount or 0) * FX.get(ccy, 1.0)
    except Exception:
        return 0.0

# ---- 采购订单 ----
wb = openpyxl.load_workbook(F1, read_only=False, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))[1:]

orders = []
order_totals = {}   # po -> 订单总价(col16, 原币)
order_ccy = {}
for r in rows:
    po, odate, supplier, pn, desc = r[4], r[3], r[6], r[20], r[26]
    qty, unit, ccy, line = r[27], r[28], r[29], r[30]
    status, scope = r[7], r[17]
    if not po:
        continue
    ym = odate.strftime("%Y-%m") if isinstance(odate, datetime.datetime) else ""
    orders.append({
        "po": str(po).strip(),
        "ym": ym,
        "supplier": (supplier or "").strip(),
        "pn": (pn or "").strip(),
        "desc": (desc or "").strip(),
        "qty": float(qty or 0),
        "unitCny": round(cny(unit, ccy), 2),
        "lineCny": round(cny(line, ccy), 2),
        "status": status or "",
        "domestic": (scope == "国内订单"),
    })
    if po not in order_totals:
        order_totals[po] = r[16] or 0
        order_ccy[po] = ccy

approved = [o for o in orders if o["status"] == "已通过"]
sum_line = sum(o["lineCny"] for o in approved)
sum_order_total = sum(cny(order_totals[p], order_ccy[p]) for p in order_totals)
print("通过行数:", len(approved), " 按行总价sum:", round(sum_line), " 按订单总价sum:", round(sum_order_total))

# ---- 应付 ----
wb2 = openpyxl.load_workbook(F2, read_only=False, data_only=True)
ws2 = wb2[wb2.sheetnames[0]]
aprows = list(ws2.iter_rows(values_only=True))[1:]

def money(v):
    if v is None: return 0.0
    s = str(v).replace("¥", "").replace("$", "").replace(",", "").strip()
    try: return float(s)
    except: return 0.0

payable = paid = 0.0
pending = []
for r in aprows:
    payable += money(r[9])
    paid += money(r[10])
    try:
        qty = float(r[5] or 0)
    except Exception:
        qty = 0
    try:
        unship = float(r[19] or 0)
    except Exception:
        unship = 0
    if unship > 0:
        eta = r[17]
        eta_s = eta.strftime("%Y-%m-%d") if isinstance(eta, datetime.datetime) else (str(eta) if eta else "—")
        pending.append({
            "po": str(r[1] or "").strip(),
            "supplier": (r[3] or "").strip(),
            "pn": (r[4] or "").strip(),
            "qty": qty,
            "amt": money(r[9]),
            "eta": eta_s,
            "unship": unship,
            "status": "待发货" if abs(unship - qty) < 1e-6 else "部分到货",
        })

print("应付:", round(payable), "已付:", round(paid), "未付:", round(payable-paid))
print("在途行数:", len(pending))

out = {
    "payable": {"payable": round(payable), "paid": round(paid), "unpaid": round(payable - paid)},
    "orders": orders,
    "pending": pending,
}
with open(r"C:\Users\kirby\Doubao\chats\2026-10-19\new-chat\procurement-dashboard\data.js", "w", encoding="utf-8") as f:
    f.write("window.DATA = " + json.dumps(out, ensure_ascii=False, indent=1) + ";\n")
print("data.js written, orders:", len(orders))
